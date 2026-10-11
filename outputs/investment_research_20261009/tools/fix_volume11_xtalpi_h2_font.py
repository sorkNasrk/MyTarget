"""Repair only the last company's H2 runs, retaining exact text and ZIP audits."""
import argparse
import copy
import datetime
import hashlib
import json
import shutil
import uuid
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree
from fix_final_report_typography import text_snapshot, text_hash, run_details

ROOT = Path(__file__).resolve().parents[1]
FILE = ROOT / 'deliverables/11_公司深研_初版.docx'
AUDIT = ROOT / 'qa/volume11_xtalpi_h2_font_fix.json'
CHAPTER = '晶泰控股 02228 HK'
THEMES = ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme']


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def zip_entries(path):
    with zipfile.ZipFile(path) as package:
        return {name: package.read(name) for name in package.namelist()}


def save_document_part_only(document, source_package, destination):
    """Retain every original non-document ZIP payload, avoiding library reserialization."""
    document_xml = etree.tostring(document._element, encoding='UTF-8', xml_declaration=True, standalone=True)
    with zipfile.ZipFile(source_package) as source, zipfile.ZipFile(destination, 'w') as output:
        for info in source.infolist():
            output.writestr(info, document_xml if info.filename == 'word/document.xml' else source.read(info.filename))


def targets(document):
    starts = [i for i, p in enumerate(document.paragraphs)
              if p.style.name == 'Heading 1' and p.text == CHAPTER]
    assert len(starts) == 1, starts
    start = starts[0]
    assert not any(p.style.name == 'Heading 1' for p in document.paragraphs[start + 1:])
    headings = [(i, p) for i, p in enumerate(document.paragraphs)
                if i > start and p.style.name == 'Heading 2']
    assert len(headings) == 23, len(headings)
    return start, headings


def verify_run(run):
    fonts = run._r.rPr.rFonts
    assert fonts.get(qn('w:ascii')) == 'Microsoft YaHei'
    assert fonts.get(qn('w:hAnsi')) == 'Microsoft YaHei'
    assert fonts.get(qn('w:eastAsia')) == '微软雅黑'
    assert all(fonts.get(qn('w:' + theme)) is None for theme in THEMES)
    assert run.font.size == Pt(13) and run.font.bold is True


def masked_xml(document, indices):
    body = copy.deepcopy(document._element.body)
    paras = body.findall(qn('w:p'))
    for index in indices:
        for run in paras[index].iter(qn('w:r')):
            rpr = run.find(qn('w:rPr'))
            if rpr is not None:
                run.remove(rpr)
    return etree.tostring(body, method='c14n')


def apply():
    assert not AUDIT.exists(), 'A completed edit must not be repeated'
    assert Path(r'C:\Windows\Fonts\msyh.ttc').is_file()
    assert Path(r'C:\Windows\Fonts\msyhbd.ttc').is_file()
    stamp = datetime.datetime.now(datetime.timezone.utc).strftime('%Y%m%dT%H%M%SZ')
    backup_dir = ROOT / 'qa' / ('backup_volume11_xtalpi_h2_' + stamp + '_' + uuid.uuid4().hex[:8])
    backup_dir.mkdir()
    backup = backup_dir / FILE.name
    shutil.copy2(FILE, backup)
    assert sha(FILE) == sha(backup)
    before_hash = sha(FILE)
    before_entries = zip_entries(FILE)
    document = Document(FILE)
    start, headings = targets(document)
    assert document.styles['Heading 2'].font.size == Pt(13)
    assert document.styles['Heading 2'].font.bold is True
    before_text = text_snapshot(document)
    before_paragraph_xml = [p._p.xml for p in document.paragraphs]
    before_styles = {s.style_id: s._element.xml for s in document.styles}
    target_indices = [i for i, _ in headings]
    before_masked = masked_xml(document, target_indices)
    affected = []
    for index, paragraph in headings:
        assert len(list(paragraph._p.iter(qn('w:r')))) == len(paragraph.runs)
        assert paragraph.runs and paragraph.text
        affected.append({'index': index, 'text': paragraph.text, 'beforeRuns': run_details(paragraph)})
        for run in paragraph.runs:
            run.font.name = 'Microsoft YaHei'
            run.font.size = Pt(13)
            run.font.bold = True
            fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn('w:ascii'), 'Microsoft YaHei')
            fonts.set(qn('w:hAnsi'), 'Microsoft YaHei')
            fonts.set(qn('w:eastAsia'), '微软雅黑')
            for theme in THEMES:
                fonts.attrib.pop(qn('w:' + theme), None)
            verify_run(run)
    changed = [i for i, p in enumerate(document.paragraphs) if p._p.xml != before_paragraph_xml[i]]
    assert changed == target_indices, (changed, target_indices)
    assert before_styles == {s.style_id: s._element.xml for s in document.styles}
    assert masked_xml(document, target_indices) == before_masked
    assert text_snapshot(document) == before_text
    save_document_part_only(document, backup, FILE)
    saved = Document(FILE)
    assert text_snapshot(saved) == before_text
    assert before_styles == {s.style_id: s._element.xml for s in saved.styles}
    assert masked_xml(saved, target_indices) == before_masked
    after_entries = zip_entries(FILE)
    assert set(after_entries) == set(before_entries)
    changed_entries = [name for name in before_entries if before_entries[name] != after_entries[name]]
    assert changed_entries == ['word/document.xml'], changed_entries
    for item in affected:
        paragraph = saved.paragraphs[item['index']]
        assert paragraph.text == item['text']
        for run in paragraph.runs:
            verify_run(run)
        item['afterRuns'] = run_details(paragraph)
    record = {
        'createdAt': datetime.datetime.now(datetime.timezone.utc).isoformat(),
        'file': str(FILE), 'backup': str(backup), 'beforeSHA256': before_hash,
        'afterTypographySHA256': sha(FILE),
        'beforeTextSHA256': text_hash(before_text),
        'afterTextSHA256': text_hash(text_snapshot(saved)),
        'passed': True, 'paragraphAndTableTextUnchanged': True,
        'globalStylesUnchanged': True, 'allOtherZIPEntriesUnchanged': True,
        'relationshipsAndURLsUnchanged': True,
        'bodyXMLUnchangedExceptTargetRunProperties': True,
        'changedZIPEntriesOnly': changed_entries,
        'changedParagraphsOnly': target_indices,
        'chapterStartIndex': start, 'chapterHeading': CHAPTER,
        'heading2StyleXMLUnchanged': saved.styles['Heading 2']._element.xml,
        'affectedParagraphs': affected,
        'confirmedOriginalProblemPages': [154, 155, 156, 157],
        'scope': 'Only all 23 H2 paragraphs in volume 11 last company 晶泰控股; explicit Microsoft YaHei / 微软雅黑 bold 13pt and clear run theme fonts. No global style, research, body/table text, URLs, or other package entries changed.',
    }
    AUDIT.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: record[k] for k in ['passed', 'backup', 'beforeSHA256', 'afterTypographySHA256', 'beforeTextSHA256', 'afterTextSHA256', 'changedParagraphsOnly', 'changedZIPEntriesOnly']}, ensure_ascii=False), flush=True)


def verify_after_toc():
    record = json.loads(AUDIT.read_text(encoding='utf-8'))
    current = Document(FILE)
    original = Document(record['backup'])
    assert text_snapshot(current) == text_snapshot(original)
    _, headings = targets(current)
    assert [p.text for _, p in headings] == [item['text'] for item in record['affectedParagraphs']]
    for _, paragraph in headings:
        for run in paragraph.runs:
            verify_run(run)
    record['afterTocSHA256'] = sha(FILE)
    record['afterTocTextSHA256'] = text_hash(text_snapshot(current))
    record['afterTocParagraphAndTableTextUnchanged'] = True
    record['afterTocTargetFontsVerified'] = True
    record['verifiedAfterTocAt'] = datetime.datetime.now(datetime.timezone.utc).isoformat()
    AUDIT.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: record[k] for k in ['afterTocSHA256', 'afterTocTextSHA256', 'afterTocParagraphAndTableTextUnchanged', 'afterTocTargetFontsVerified']}, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--verify-after-toc', action='store_true')
    args = parser.parse_args()
    verify_after_toc() if args.verify_after_toc else apply()
