"""Apply local typography repairs with exact text and recoverable-file audits.

No report wording, links, table values, economics, or global styles change.
"""
import argparse
import datetime
import hashlib
import json
import shutil
import uuid
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
AUDIT = ROOT / 'qa/final_report_typography_fix.json'
SOURCE_IDS = {
    9: [
        '2026082801141', '2026033002160', '2026100700513',
        '2026010601500', '2026072300835', '2026082000533',
        '2026092300616', '2026060402931',
    ],
    10: [
        '2026041000366', '2026031200310', '2026090800186',
        '2026082600506', '2026082600316', '2026052800509',
        '2026100100037', '2026090100613', '2026080200003',
        '2026032401608', '2026100700656',
    ],
    11: [
        '2026091600294', '2026042201748', '2026100601007',
        '2026072800183', '2026092300318', '2025060100117',
        '2025012200037',
    ],
}
HEADINGS = [
    '十三 外资目标价 治理及核验边界',
    '半年经营逻辑与失效路径',
    '一年经营逻辑与失效路径',
    '来源与核验记录 不计入独立正文字数',
    '四期限逻辑与失效路径',
]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def text_snapshot(document):
    """Each ordinary paragraph and table-cell paragraph, including nested tables."""
    return {
        'paragraphs': [paragraph.text for paragraph in document.paragraphs],
        'tableParagraphs': [
            ''.join(text_node.text or '' for text_node in node.iter(qn('w:t')))
            for table in document.tables
            for node in table._tbl.iter(qn('w:p'))
        ],
    }


def text_hash(snapshot):
    return hashlib.sha256(
        json.dumps(snapshot, ensure_ascii=False, separators=(',', ':')).encode('utf-8')
    ).hexdigest()


def run_details(paragraph):
    return [
        {
            'text': run.text,
            'font': run.font.name,
            'sizePt': run.font.size.pt if run.font.size else None,
            'bold': run.font.bold,
            'xml': run._r.xml,
        }
        for run in paragraph.runs
    ]


def sources(document, identifiers):
    matched = []
    for identifier in identifiers:
        paragraphs = [
            (index, paragraph)
            for index, paragraph in enumerate(document.paragraphs)
            if identifier in paragraph.text
        ]
        assert len(paragraphs) == 1, (identifier, len(paragraphs))
        index, paragraph = paragraphs[0]
        assert paragraph.text.startswith('- '), paragraph.text
        assert 'https://' in paragraph.text, paragraph.text
        assert paragraph.style.name == 'Normal', paragraph.style.name
        matched.append((identifier, index, paragraph))
    assert len({index for _, index, _ in matched}) == len(identifiers)
    return matched


def apply():
    assert not AUDIT.exists(), 'Do not repeat a completed typography edit batch'
    stamp = datetime.datetime.now(datetime.timezone.utc).strftime('%Y%m%dT%H%M%SZ')
    backup_dir = ROOT / 'qa' / ('backup_typography_' + stamp + '_' + uuid.uuid4().hex[:8])
    backup_dir.mkdir()
    files = [ROOT / 'deliverables' / f'{index:02d}_公司深研_初版.docx' for index in range(9, 13)]
    originals = {}
    for file in files:
        backup = backup_dir / file.name
        shutil.copy2(file, backup)
        assert sha(file) == sha(backup)
        originals[file] = backup

    records = []
    for index, file in zip(range(9, 13), files):
        document = Document(file)
        before_text = text_snapshot(document)
        before_hash = sha(file)
        before_xml = [paragraph._p.xml for paragraph in document.paragraphs]
        affected = []
        if index in SOURCE_IDS:
            for identifier, paragraph_index, paragraph in sources(document, SOURCE_IDS[index]):
                affected.append({
                    'index': paragraph_index, 'text': paragraph.text,
                    'sourceId': identifier, 'beforeRuns': run_details(paragraph),
                    'change': 'Only source paragraph runs set to 9.5pt; text and URL unchanged',
                })
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
        else:
            assert Path(r'C:\Windows\Fonts\msyh.ttc').is_file()
            assert Path(r'C:\Windows\Fonts\msyhbd.ttc').is_file()
            for heading in HEADINGS:
                matched = [(i, p) for i, p in enumerate(document.paragraphs) if p.text == heading]
                assert len(matched) == 1, (heading, len(matched))
                paragraph_index, paragraph = matched[0]
                assert paragraph.style.name == 'Heading 2'
                assert paragraph.style.font.size == Pt(13)
                assert paragraph.style.font.bold is True
                affected.append({
                    'index': paragraph_index, 'text': paragraph.text,
                    'beforeRuns': run_details(paragraph),
                    'change': 'Only target H2 runs set Microsoft YaHei / 微软雅黑, bold 13pt; theme font attributes cleared',
                })
                for run in paragraph.runs:
                    run.font.name = 'Microsoft YaHei'
                    run.font.size = Pt(13)
                    run.font.bold = True
                    fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
                    fonts.set(qn('w:ascii'), 'Microsoft YaHei')
                    fonts.set(qn('w:hAnsi'), 'Microsoft YaHei')
                    fonts.set(qn('w:eastAsia'), '微软雅黑')
                    for theme in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme']:
                        fonts.attrib.pop(qn('w:' + theme), None)

        expected_indices = {item['index'] for item in affected}
        changed_indices = {i for i, paragraph in enumerate(document.paragraphs) if paragraph._p.xml != before_xml[i]}
        assert changed_indices == expected_indices, (changed_indices, expected_indices)
        assert text_snapshot(document) == before_text
        document.save(file)
        saved = Document(file)
        assert text_snapshot(saved) == before_text
        for item in affected:
            paragraph = saved.paragraphs[item['index']]
            item['afterRuns'] = run_details(paragraph)
            assert paragraph.text == item['text']
            if index in SOURCE_IDS:
                assert all(run.font.size == Pt(9.5) for run in paragraph.runs)
            else:
                for run in paragraph.runs:
                    fonts = run._r.rPr.rFonts
                    assert fonts.get(qn('w:ascii')) == 'Microsoft YaHei'
                    assert fonts.get(qn('w:hAnsi')) == 'Microsoft YaHei'
                    assert fonts.get(qn('w:eastAsia')) == '微软雅黑'
                    assert all(fonts.get(qn('w:' + theme)) is None for theme in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme'])
                    assert run.font.size == Pt(13) and run.font.bold is True
        records.append({
            'index': index, 'file': str(file), 'backup': str(originals[file]),
            'beforeSHA256': before_hash, 'afterTypographySHA256': sha(file),
            'beforeTextSHA256': text_hash(before_text), 'afterTextSHA256': text_hash(text_snapshot(saved)),
            'paragraphAndTableTextUnchanged': True,
            'changedParagraphsOnly': sorted(changed_indices), 'affectedParagraphs': affected,
        })
    audit = {
        'createdAt': datetime.datetime.now(datetime.timezone.utc).isoformat(),
        'backupDirectory': str(backup_dir), 'records': records, 'passed': True,
        'annualReportMatcherCorrection': 'Volume 11 2025 annual report is 2026042201748, as confirmed from current paragraph and original p102 PNG; no URL changed.',
        'scope': 'Typography only: 8 volume 9 sources, 11 volume 10 sources, 7 volume 11 sources, 5 volume 12 H2 headings.',
    }
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({
        'audit': str(AUDIT), 'backup': str(backup_dir), 'passed': True,
        'records': [{key: record[key] for key in ['index', 'beforeSHA256', 'afterTypographySHA256', 'changedParagraphsOnly']} for record in records],
    }, ensure_ascii=False), flush=True)


def verify_after_toc():
    audit = json.loads(AUDIT.read_text(encoding='utf-8'))
    for record in audit['records']:
        current = Document(record['file'])
        original = Document(record['backup'])
        assert text_snapshot(current) == text_snapshot(original), record['file']
        record['afterTocSHA256'] = sha(Path(record['file']))
        record['afterTocTextSHA256'] = text_hash(text_snapshot(current))
        record['afterTocParagraphAndTableTextUnchanged'] = True
        record['afterTocAffectedParagraphs'] = []
        for item in record['affectedParagraphs']:
            matched = [p for p in current.paragraphs if p.text == item['text']]
            assert len(matched) == 1, item['text']
            record['afterTocAffectedParagraphs'].append({'text': item['text'], 'runs': run_details(matched[0])})
    audit['verifiedAfterTocAt'] = datetime.datetime.now(datetime.timezone.utc).isoformat()
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'verifiedAfterToc': True, 'textUnchanged': True, 'files': [record['file'] for record in audit['records']]}, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--verify-after-toc', action='store_true')
    args = parser.parse_args()
    if args.verify_after_toc:
        verify_after_toc()
    else:
        apply()
