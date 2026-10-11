"""Isolate exact source-paragraph wrap settings in a disposable QA DOCX."""
import copy
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
source = Document(ROOT / 'deliverables/10_公司深研_初版.docx')
target = next(p for p in source.paragraphs if '2026041000366' in p.text)
document = Document()
for attr in ['page_width', 'page_height', 'left_margin', 'right_margin', 'top_margin', 'bottom_margin']:
    setattr(document.sections[0], attr, getattr(source.sections[0], attr))
normal = document.styles['Normal']
normal.font.name = '宋体'
normal.font.size = Pt(11)
normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
normal.paragraph_format.line_spacing = 1.4

for label, flag, arial, language in [
    ('A character wrapping', True, False, False),
    ('B compact Latin URL', False, True, False),
    ('C zh CN language', False, False, True),
    ('D wrapping and language', True, False, True),
]:
    document.add_paragraph(label)
    paragraph = copy.deepcopy(target._p)
    document._body._body.insert(-1, paragraph)
    if flag:
        wrap = OxmlElement('w:wordWrap')
        wrap.set(qn('w:val'), '0')
        paragraph.get_or_add_pPr().append(wrap)
    for run in paragraph.iter(qn('w:r')):
        if arial and 'https://' in ''.join(run.itertext()):
            fonts = run.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn('w:ascii'), 'Arial')
            fonts.set(qn('w:hAnsi'), 'Arial')
        if language:
            lang = OxmlElement('w:lang')
            lang.set(qn('w:val'), 'zh-CN')
            lang.set(qn('w:eastAsia'), 'zh-CN')
            run.get_or_add_rPr().append(lang)
out = ROOT / 'qa/source_wrap_probe.docx'
document.save(out)
print(json.dumps({'probe': str(out)}, ensure_ascii=False))
