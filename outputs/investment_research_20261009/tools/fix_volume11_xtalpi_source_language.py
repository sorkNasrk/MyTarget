"""Fix six observed 晶泰 source-label wraps with Chinese run language only."""
import argparse
import datetime
import json
import shutil
import uuid
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fix_final_report_typography import text_snapshot, text_hash, run_details
from fix_volume11_xtalpi_h2_font import sha, zip_entries, masked_xml, save_document_part_only

ROOT = Path(__file__).resolve().parents[1]
FILE = ROOT / 'deliverables/11_公司深研_初版.docx'
AUDIT = ROOT / 'qa/volume11_xtalpi_source_language_fix.json'
IDS = ['2026081900919', '2026100601423', '2026052701684', '2026091701791', '2026100203343', '2026100502035']


def matches(document):
    found = []
    for identifier in IDS:
        result = [(i, p) for i, p in enumerate(document.paragraphs) if identifier in p.text]
        assert len(result) == 1, identifier
        index, paragraph = result[0]
        assert index > 953 and paragraph.style.name == 'Normal' and paragraph.text.startswith('- ')
        found.append((identifier, index, paragraph))
    return found


def verify_language(paragraph):
    for run in paragraph.runs:
        lang = run._r.rPr.find(qn('w:lang'))
        assert lang.get(qn('w:val')) == 'zh-CN' and lang.get(qn('w:eastAsia')) == 'zh-CN'


def apply():
    assert not AUDIT.exists()
    backup_dir = ROOT / 'qa' / ('backup_volume11_xtalpi_source_' + datetime.datetime.now(datetime.timezone.utc).strftime('%Y%m%dT%H%M%SZ') + '_' + uuid.uuid4().hex[:8])
    backup_dir.mkdir()
    backup = backup_dir / FILE.name
    shutil.copy2(FILE, backup)
    assert sha(backup) == sha(FILE)
    before_hash = sha(FILE)
    original_entries = zip_entries(FILE)
    document = Document(FILE)
    before_text = text_snapshot(document)
    before_styles = document.styles.element.xml
    before_paragraphs = [p._p.xml for p in document.paragraphs]
    found = matches(document)
    indices = [i for _, i, _ in found]
    before_masked = masked_xml(document, indices)
    affected = []
    for identifier, index, paragraph in found:
        before_sizes = [r.font.size for r in paragraph.runs]
        affected.append({'sourceId': identifier, 'index': index, 'text': paragraph.text, 'beforeRuns': run_details(paragraph)})
        for run in paragraph.runs:
            rpr = run._r.get_or_add_rPr()
            lang = rpr.find(qn('w:lang'))
            if lang is None:
                lang = OxmlElement('w:lang')
                rpr.append(lang)
            lang.set(qn('w:val'), 'zh-CN')
            lang.set(qn('w:eastAsia'), 'zh-CN')
        assert [r.font.size for r in paragraph.runs] == before_sizes
    changed = [i for i, p in enumerate(document.paragraphs) if p._p.xml != before_paragraphs[i]]
    assert changed == indices
    assert document.styles.element.xml == before_styles
    assert masked_xml(document, indices) == before_masked
    assert text_snapshot(document) == before_text
    save_document_part_only(document, backup, FILE)
    saved = Document(FILE)
    assert text_snapshot(saved) == before_text
    assert saved.styles.element.xml == before_styles
    assert masked_xml(saved, indices) == before_masked
    final_entries = zip_entries(FILE)
    changed_entries = [name for name in original_entries if final_entries[name] != original_entries[name]]
    assert changed_entries == ['word/document.xml']
    for item in affected:
        paragraph = saved.paragraphs[item['index']]
        verify_language(paragraph)
        item['afterRuns'] = run_details(paragraph)
    audit = {
        'createdAt': datetime.datetime.now(datetime.timezone.utc).isoformat(),
        'passed': True, 'file': str(FILE), 'backup': str(backup),
        'beforeSHA256': before_hash, 'afterTypographySHA256': sha(FILE),
        'beforeTextSHA256': text_hash(before_text), 'afterTextSHA256': text_hash(text_snapshot(saved)),
        'changedParagraphsOnly': indices, 'changedZIPEntriesOnly': changed_entries,
        'paragraphAndTableTextUnchanged': True, 'stylesAndTablesXMLUnchanged': True,
        'sourceURLsAndRelationshipsUnchanged': True, 'bodyXMLUnchangedExceptTargetRunProperties': True,
        'fontsAndSizesUnchanged': True, 'allOtherZIPEntriesUnchanged': True,
        'affectedParagraphs': affected,
        'confirmedOriginalPage': 165,
        'scope': 'Only six source paragraphs whose labels visibly split final Chinese characters at line ends; run language zh-CN only. Fonts and inherited 11pt size remain exact; all text, URLs, H2 repairs and shared styles retained.',
    }
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: audit[k] for k in ['passed', 'backup', 'beforeSHA256', 'afterTypographySHA256', 'beforeTextSHA256', 'afterTextSHA256', 'changedParagraphsOnly']}, ensure_ascii=False), flush=True)


def verify_after_toc():
    audit = json.loads(AUDIT.read_text(encoding='utf-8'))
    saved, original = Document(FILE), Document(audit['backup'])
    assert text_snapshot(saved) == text_snapshot(original)
    for _, _, paragraph in matches(saved):
        verify_language(paragraph)
    audit['afterTocSHA256'] = sha(FILE)
    audit['afterTocTextSHA256'] = text_hash(text_snapshot(saved))
    audit['afterTocTextAndURLsUnchanged'] = True
    audit['afterTocLanguageVerified'] = True
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: audit[k] for k in ['afterTocSHA256', 'afterTocTextSHA256', 'afterTocTextAndURLsUnchanged', 'afterTocLanguageVerified']}, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--verify-after-toc', action='store_true')
    args = parser.parse_args()
    verify_after_toc() if args.verify_after_toc else apply()
