"""Translate two additional missing-horizon labels, preserving all other text."""
import datetime, hashlib, json
from pathlib import Path
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
changes=[(1,'中（公开新闻转述，期限null）','中（公开新闻转述，目标期限未披露）'),
         (5,'中（公开新闻转述，目标期限缺失null）','中（公开新闻转述，目标期限未披露）')]
texts=lambda d:[p.text for p in d.paragraphs]+[c.text for t in d.tables for r in t.rows for c in r.cells]
records=[]
for volume,old,new in changes:
    f=ROOT/'deliverables'/f'{volume:02d}_公司深研_初版.docx'
    d=Document(f); original=texts(d)
    cells=[c for t in d.tables for r in t.rows for c in r.cells if c.text==old]
    assert len(cells)==1
    before=hashlib.sha256(f.read_bytes()).hexdigest()
    p=cells[0].paragraphs[0]
    assert p.text==old and p.runs
    assert [r.text for r in p.runs]==[old[:-5],'null','）']
    p.runs[0].text='中（公开新闻转述，目标期限'
    p.runs[1].text='未披露'
    assert cells[0].text==new
    d.save(f)
    assert [t.replace(old,new) for t in original]==texts(Document(f))
    records.append({'volume':volume,'file':str(f),'beforeSHA256':before,'afterSHA256':hashlib.sha256(f.read_bytes()).hexdigest(),'oldLabel':old,'newLabel':new})
q={'passed':True,'createdAt':datetime.datetime.now(datetime.timezone.utc).isoformat(),'records':records,'change':'Presentation-only translation of missing broker horizon; frozen models, source inputs and economics unchanged.'}
(ROOT/'qa/other_confidence_label_fix.json').write_text(json.dumps(q,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(q,ensure_ascii=False))
