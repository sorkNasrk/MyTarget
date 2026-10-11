"""Repair the confirmed volume 12 H2 fallback font without changing text."""
import argparse
import datetime
import hashlib
import json
import shutil
import uuid
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from fix_final_report_typography import text_snapshot, text_hash

ROOT = Path(__file__).resolve().parents[1]
FILE = ROOT / 'deliverables/12_公司深研_初版.docx'
AUDIT = ROOT / 'qa/volume12_h2_font_style_fix.json'


def sha(file):
    return hashlib.sha256(file.read_bytes()).hexdigest()


def verify_font(style):
    fonts = style._element.rPr.rFonts
    assert fonts.get(qn('w:ascii')) == 'Microsoft YaHei'
    assert fonts.get(qn('w:hAnsi')) == 'Microsoft YaHei'
    assert fonts.get(qn('w:eastAsia')) == '微软雅黑'
    assert all(fonts.get(qn('w:' + theme)) is None for theme in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme'])
    assert style.font.size == Pt(13) and style.font.bold is True


def apply():
    assert not AUDIT.exists()
    stamp = datetime.datetime.now(datetime.timezone.utc).strftime('%Y%m%dT%H%M%SZ')
    backup_dir = ROOT / 'qa' / ('backup_h2style_' + stamp + '_' + uuid.uuid4().hex[:8])
    backup_dir.mkdir()
    backup = backup_dir / FILE.name
    shutil.copy2(FILE, backup)
    before_hash = sha(FILE)
    assert before_hash == sha(backup)
    document = Document(FILE)
    before_text = text_snapshot(document)
    before_body_xml = document._element.body.xml
    before_styles = {style.style_id: style._element.xml for style in document.styles}
    style = document.styles['Heading 2']
    assert style.font.size == Pt(13) and style.font.bold is True
    before_style = style._element.xml
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(13)
    style.font.bold = True
    fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:ascii'), 'Microsoft YaHei')
    fonts.set(qn('w:hAnsi'), 'Microsoft YaHei')
    fonts.set(qn('w:eastAsia'), '微软雅黑')
    for theme in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme']:
        fonts.attrib.pop(qn('w:' + theme), None)
    verify_font(style)
    changed_styles = [item.style_id for item in document.styles if item._element.xml != before_styles[item.style_id]]
    assert changed_styles == ['Heading2'], changed_styles
    assert document._element.body.xml == before_body_xml
    assert text_snapshot(document) == before_text
    document.save(FILE)
    saved = Document(FILE)
    verify_font(saved.styles['Heading 2'])
    assert saved._element.body.xml == before_body_xml
    assert text_snapshot(saved) == before_text
    audit = {
        'createdAt': datetime.datetime.now(datetime.timezone.utc).isoformat(),
        'file': str(FILE), 'backup': str(backup),
        'beforeSHA256': before_hash, 'afterStyleSHA256': sha(FILE),
        'beforeTextSHA256': text_hash(before_text), 'afterTextSHA256': text_hash(text_snapshot(saved)),
        'passed': True, 'paragraphAndTableTextUnchanged': True, 'bodyXMLUnchanged': True,
        'changedStyleIdsOnly': changed_styles,
        'beforeStyleXML': before_style, 'afterStyleXML': saved.styles['Heading 2']._element.xml,
        'affectedHeading2Paragraphs': [
            {'index': index, 'text': paragraph.text}
            for index, paragraph in enumerate(saved.paragraphs)
            if paragraph.style.name == 'Heading 2'
        ],
        'confirmedAdditionalOriginalPages': [85, 86, 87, 88],
        'change': 'Only volume 12 Heading2 style font: explicit Microsoft YaHei / 微软雅黑, remove ASCII/HAnsi/EastAsia theme font attributes, preserve bold 13pt.',
    }
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'audit': str(AUDIT), 'passed': True, 'changedStyleIdsOnly': changed_styles, 'headings': len(audit['affectedHeading2Paragraphs']), 'afterStyleSHA256': audit['afterStyleSHA256']}, ensure_ascii=False), flush=True)


def verify_after_toc():
    audit = json.loads(AUDIT.read_text(encoding='utf-8'))
    current = Document(FILE)
    before = Document(audit['backup'])
    assert text_snapshot(current) == text_snapshot(before)
    verify_font(current.styles['Heading 2'])
    audit['afterTocSHA256'] = sha(FILE)
    audit['afterTocParagraphAndTableTextUnchanged'] = True
    audit['afterTocStyleXML'] = current.styles['Heading 2']._element.xml
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'afterTocVerified': True, 'afterTocSHA256': audit['afterTocSHA256']}, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--verify-after-toc', action='store_true')
    args = parser.parse_args()
    verify_after_toc() if args.verify_after_toc else apply()
