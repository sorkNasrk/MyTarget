"""Repair remaining local source wraps using Chinese-language line breaking.

Only two volume-9 sources, eleven volume-10 sources, and the confirmed
volume-9 boundary heading receive direct typography attributes. Exact visible
paragraph/table text, URLs, global styles and package relationships are audited.
"""
import argparse
import datetime
import hashlib
import json
import re
import shutil
import uuid
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

from fix_final_report_typography import text_snapshot, text_hash, run_details

ROOT = Path(__file__).resolve().parents[1]
AUDIT = ROOT / 'qa/final_source_language_fix.json'
IDS = {
    9: ['2026082801141', '2026010601500'],
    10: [
        '2026041000366', '2026031200310', '2026090800186',
        '2026082600506', '2026082600316', '2026052800509',
        '2026100100037', '2026090100613', '2026080200003',
        '2026032401608', '2026100700656',
    ],
}


def sha(file):
    return hashlib.sha256(Path(file).read_bytes()).hexdigest()


def utc():
    return datetime.datetime.now(datetime.timezone.utc).isoformat()


def all_xml_visible_text(file):
    with zipfile.ZipFile(file) as package:
        return {
            name: [node.text or '' for node in etree.fromstring(package.read(name)).iter(qn('w:t'))]
            for name in package.namelist()
            if name.startswith('word/') and name.endswith('.xml')
        }


def all_relationship_parts(file):
    with zipfile.ZipFile(file) as package:
        return {name: package.read(name) for name in package.namelist() if name.endswith('.rels')}


def source_urls(document):
    return re.findall(r'https?://[^\s)]+', '\n'.join(p.text for p in document.paragraphs))


def attributes_verified(paragraph, heading=False):
    for run in paragraph.runs:
        if heading:
            fonts = run._r.rPr.rFonts
            assert fonts.get(qn('w:ascii')) == 'Microsoft YaHei'
            assert fonts.get(qn('w:hAnsi')) == 'Microsoft YaHei'
            assert fonts.get(qn('w:eastAsia')) == '微软雅黑'
            assert run.font.size == Pt(13) and run.font.bold is True
        else:
            language = run._r.rPr.find(qn('w:lang'))
            assert language is not None
            assert language.get(qn('w:val')) == 'zh-CN'
            assert language.get(qn('w:eastAsia')) == 'zh-CN'
            assert run.font.size == Pt(9.5)


def apply():
    assert not AUDIT.exists(), 'Completed repair batch must not be repeated'
    backup_dir = ROOT / 'qa' / (
        'backup_source_language_' + datetime.datetime.now(datetime.timezone.utc).strftime('%Y%m%dT%H%M%SZ')
        + '_' + uuid.uuid4().hex[:8]
    )
    backup_dir.mkdir()
    records = []
    untouched = {
        str(ROOT / 'deliverables' / f'{n:02d}_公司深研_初版.docx'):
        sha(ROOT / 'deliverables' / f'{n:02d}_公司深研_初版.docx')
        for n in [11, 12]
    }
    for volume in [9, 10]:
        file = ROOT / 'deliverables' / f'{volume:02d}_公司深研_初版.docx'
        backup = backup_dir / file.name
        shutil.copy2(file, backup)
        before_hash = sha(file)
        assert sha(backup) == before_hash
        document = Document(file)
        before = text_snapshot(document)
        visible_xml = all_xml_visible_text(file)
        urls = source_urls(document)
        before_xml = [p._p.xml for p in document.paragraphs]
        before_styles = document.styles.element.xml
        before_tables = [t._tbl.xml for t in document.tables]
        affected = []
        for identifier in IDS[volume]:
            matched = [(i, p) for i, p in enumerate(document.paragraphs) if identifier in p.text]
            assert len(matched) == 1, (identifier, len(matched))
            index, paragraph = matched[0]
            assert paragraph.style.name == 'Normal' and paragraph.text.startswith('- ')
            assert all(run.font.size == Pt(9.5) for run in paragraph.runs)
            item = {'index': index, 'text': paragraph.text, 'sourceId': identifier, 'heading': False,
                    'beforeRuns': run_details(paragraph)}
            for run in paragraph.runs:
                properties = run._r.get_or_add_rPr()
                language = properties.find(qn('w:lang'))
                if language is None:
                    language = OxmlElement('w:lang')
                    properties.append(language)
                language.set(qn('w:val'), 'zh-CN')
                language.set(qn('w:eastAsia'), 'zh-CN')
            item['change'] = 'Only run language set to zh-CN for Chinese URL line breaking; 9.5pt, text and URL retained'
            affected.append(item)
        if volume == 9:
            matched = [(i, p) for i, p in enumerate(document.paragraphs) if p.text == '核验边界']
            assert len(matched) == 1
            index, paragraph = matched[0]
            assert paragraph.style.name == 'Heading 2'
            item = {'index': index, 'text': paragraph.text, 'heading': True,
                    'beforeRuns': run_details(paragraph)}
            for run in paragraph.runs:
                run.font.name = 'Microsoft YaHei'
                run.font.size = Pt(13)
                run.font.bold = True
                fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
                fonts.set(qn('w:ascii'), 'Microsoft YaHei')
                fonts.set(qn('w:hAnsi'), 'Microsoft YaHei')
                fonts.set(qn('w:eastAsia'), '微软雅黑')
                for theme in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme']:
                    fonts.attrib.pop(qn('w:' + theme), None)
            item['change'] = 'Only confirmed 核验边界 H2 run font explicitly YaHei bold13pt; shared style unchanged'
            affected.append(item)
        changed = [i for i, p in enumerate(document.paragraphs) if p._p.xml != before_xml[i]]
        assert set(changed) == {item['index'] for item in affected}
        assert document.styles.element.xml == before_styles
        assert [t._tbl.xml for t in document.tables] == before_tables
        assert text_snapshot(document) == before
        document.save(file)
        saved = Document(file)
        assert text_snapshot(saved) == before
        assert all_xml_visible_text(file) == visible_xml
        assert source_urls(saved) == urls
        for item in affected:
            paragraph = saved.paragraphs[item['index']]
            attributes_verified(paragraph, item['heading'])
            item['afterRuns'] = run_details(paragraph)
        records.append({
            'volume': volume, 'file': str(file), 'backup': str(backup),
            'beforeSHA256': before_hash, 'afterTypographySHA256': sha(file),
            'beforeTextSHA256': text_hash(before), 'afterTextSHA256': text_hash(text_snapshot(saved)),
            'paragraphAndTableTextUnchanged': True, 'allWordXMLVisibleTextUnchanged': True,
            'sourceURLsUnchanged': True, 'sourceURLs': urls, 'stylesAndTablesXMLUnchanged': True,
            'changedParagraphsOnly': changed, 'affectedParagraphs': affected,
        })
    audit = {'createdAt': utc(), 'passed': True, 'backupDirectory': str(backup_dir), 'records': records,
             'untouchedVolumesAtEdit': untouched,
             'probe': str(ROOT / 'qa/render_source_wrap_probe/page-1.png'),
             'scope': 'Formatting only: 13 source paragraphs and one volume9 boundary heading; no manuscript/model edits'}
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'passed': True, 'audit': str(AUDIT), 'backup': str(backup_dir),
                      'changed': [{'volume': r['volume'], 'paragraphs': r['changedParagraphsOnly']} for r in records]}, ensure_ascii=False), flush=True)


def verify(tag=None):
    audit = json.loads(AUDIT.read_text(encoding='utf-8'))
    manifest = json.loads((ROOT / 'qa' / f'report_render_manifest_{tag}.json').read_text(encoding='utf-8')) if tag else None
    for record in audit['records']:
        document = Document(record['file'])
        original = Document(record['backup'])
        assert text_snapshot(document) == text_snapshot(original)
        assert all_xml_visible_text(record['file']) == all_xml_visible_text(record['backup'])
        assert source_urls(document) == source_urls(original)
        assert all_relationship_parts(record['file']) == all_relationship_parts(record['backup'])
        for item in record['affectedParagraphs']:
            matches = [p for p in document.paragraphs if p.text == item['text']]
            assert len(matches) == 1
            attributes_verified(matches[0], item['heading'])
        record['currentDocxSHA256'] = sha(record['file'])
        record['afterTocTextUnchanged'] = True
        record['afterTocURLsUnchanged'] = True
        if manifest:
            rendered = next(r for r in manifest['records'] if r['index'] == record['volume'])
            assert rendered['docxSHA256'] == sha(record['file'])
            assert rendered['pdfSHA256'] == sha(rendered['pdf'])
            record['pdf'] = rendered['pdf']
            record['currentPdfSHA256'] = rendered['pdfSHA256']
            record['pages'] = rendered['pages']
            from pypdf import PdfReader
            pdf_text = re.sub(r'\s+', '', '\n'.join(page.extract_text() or '' for page in PdfReader(rendered['pdf']).pages))
            affected_urls = [url for item in record['affectedParagraphs'] if not item['heading']
                             for url in re.findall(r'https?://[^\s)]+', item['text'])]
            assert all(re.sub(r'\s+', '', url) in pdf_text for url in affected_urls)
            record['affectedSourceURLsPresentInRenderedPdf'] = True
            record['verifiedAffectedURLs'] = affected_urls
        record['allRelationshipPartsUnchanged'] = True
    audit['verifiedAt'] = utc()
    if tag:
        audit['verifiedRenderTag'] = tag
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'passed': True, 'verifiedTag': tag, 'audit': str(AUDIT),
                      'records': [{key: r[key] for key in ['volume', 'currentDocxSHA256'] + (['currentPdfSHA256', 'pages'] if tag else [])} for r in audit['records']]}, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--verify', action='store_true')
    parser.add_argument('--tag')
    args = parser.parse_args()
    verify(args.tag) if args.verify else apply()
