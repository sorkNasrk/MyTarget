"""Repair one confirmed overprinted volume9 tail H2 with direct YaHei font."""
import argparse
import datetime
import json
import shutil
import uuid
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from fix_final_report_typography import run_details, text_snapshot, text_hash
from fix_final_source_language import all_xml_visible_text, source_urls, sha

ROOT = Path(__file__).resolve().parents[1]
FILE = ROOT / 'deliverables/09_公司深研_初版.docx'
AUDIT = ROOT / 'qa/volume9_tail_heading_fix.json'
TEXT = '十年经营持续性展望'


def validate(paragraph):
    assert paragraph.style.name == 'Heading 2'
    for run in paragraph.runs:
        fonts = run._r.rPr.rFonts
        assert fonts.get(qn('w:ascii')) == 'Microsoft YaHei'
        assert fonts.get(qn('w:hAnsi')) == 'Microsoft YaHei'
        assert fonts.get(qn('w:eastAsia')) == '微软雅黑'
        assert run.font.size == Pt(13) and run.font.bold is True


def apply():
    assert not AUDIT.exists()
    backup_dir = ROOT / 'qa' / (
        'backup_volume9_tail_h2_' + datetime.datetime.now(datetime.timezone.utc).strftime('%Y%m%dT%H%M%SZ')
        + '_' + uuid.uuid4().hex[:8]
    )
    backup_dir.mkdir()
    backup = backup_dir / FILE.name
    shutil.copy2(FILE, backup)
    assert sha(FILE) == sha(backup)
    manifest = ROOT / 'qa/report_render_manifest_final92source.json'
    saved_manifest = ROOT / 'qa/report_render_manifest_final92source_before_tail_h2.json'
    assert not saved_manifest.exists()
    shutil.copy2(manifest, saved_manifest)
    assert sha(manifest) == sha(saved_manifest)
    document = Document(FILE)
    before = text_snapshot(document)
    visible = all_xml_visible_text(FILE)
    urls = source_urls(document)
    before_xml = [p._p.xml for p in document.paragraphs]
    before_styles = document.styles.element.xml
    before_tables = [t._tbl.xml for t in document.tables]
    matched = [(i, p) for i, p in enumerate(document.paragraphs) if p.text == TEXT]
    assert len(matched) == 1
    index, paragraph = matched[0]
    assert paragraph.style.name == 'Heading 2'
    before_runs = run_details(paragraph)
    for run in paragraph.runs:
        run.font.name = 'Microsoft YaHei'
        run.font.size = Pt(13)
        run.font.bold = True
        fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn('w:ascii'), 'Microsoft YaHei')
        fonts.set(qn('w:hAnsi'), 'Microsoft YaHei')
        fonts.set(qn('w:eastAsia'), '微软雅黑')
        for theme in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme']:
            fonts.attrib.pop(qn('w:' + theme), None)
    changed = [i for i, p in enumerate(document.paragraphs) if p._p.xml != before_xml[i]]
    assert changed == [index]
    assert document.styles.element.xml == before_styles
    assert [t._tbl.xml for t in document.tables] == before_tables
    assert text_snapshot(document) == before
    document.save(FILE)
    saved = Document(FILE)
    assert text_snapshot(saved) == before
    assert all_xml_visible_text(FILE) == visible
    assert source_urls(saved) == urls
    validate(saved.paragraphs[index])
    audit = {
        'createdAt': datetime.datetime.now(datetime.timezone.utc).isoformat(),
        'file': str(FILE), 'backup': str(backup), 'passed': True,
        'beforeSHA256': sha(backup), 'afterTypographySHA256': sha(FILE),
        'text': TEXT, 'changedParagraphsOnly': changed,
        'beforeTextSHA256': text_hash(before), 'afterTextSHA256': text_hash(text_snapshot(saved)),
        'paragraphAndTableTextUnchanged': True, 'allWordXMLVisibleTextUnchanged': True,
        'sourceURLsUnchanged': True, 'stylesAndTablesXMLUnchanged': True,
        'beforeRuns': before_runs, 'afterRuns': run_details(saved.paragraphs[index]),
        'preservedIntermediateManifest': str(saved_manifest),
        'originalConfirmedPage': 138,
        'change': 'Only confirmed 十年经营持续性展望 H2 run explicit YaHei13bold; shared style and content retained',
    }
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'passed': True, 'audit': str(AUDIT), 'backup': str(backup), 'changedParagraphsOnly': changed}, ensure_ascii=False), flush=True)


def verify():
    audit = json.loads(AUDIT.read_text(encoding='utf-8'))
    document = Document(FILE)
    original = Document(audit['backup'])
    assert text_snapshot(document) == text_snapshot(original)
    assert all_xml_visible_text(FILE) == all_xml_visible_text(audit['backup'])
    assert source_urls(document) == source_urls(original)
    matched = [p for p in document.paragraphs if p.text == TEXT]
    assert len(matched) == 1
    validate(matched[0])
    manifest = json.loads((ROOT / 'qa/report_render_manifest_final92source.json').read_text(encoding='utf-8'))
    record = next(r for r in manifest['records'] if r['index'] == 9)
    assert sha(FILE) == record['docxSHA256']
    assert sha(record['pdf']) == record['pdfSHA256']
    audit.update({'verifiedAt': datetime.datetime.now(datetime.timezone.utc).isoformat(),
                  'currentDocxSHA256': sha(FILE), 'currentPdfSHA256': record['pdfSHA256'],
                  'pdf': record['pdf'], 'pages': record['pages'], 'afterTocExactTextAndURLsUnchanged': True})
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({key: audit[key] for key in ['passed', 'currentDocxSHA256', 'currentPdfSHA256', 'pages']}, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--verify', action='store_true')
    args = parser.parse_args()
    verify() if args.verify else apply()
