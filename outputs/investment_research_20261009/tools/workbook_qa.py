import json, math
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
B=json.loads((ROOT/'data/research_bundle.json').read_text(encoding='utf-8'))
file=ROOT/'deliverables/证券研究总表_初版.xlsx'
values=load_workbook(file,data_only=True,read_only=True)
formulas=load_workbook(file,data_only=False,read_only=True)
checks=[];errors=[]
for s in values:
    for row in s:
        for c in row:
            if c.data_type=='e':errors.append({'sheet':s.title,'cell':c.coordinate,'error':c.value})
            if isinstance(c.value,str) and len(c.value)>32767:errors.append({'sheet':s.title,'cell':c.coordinate,'error':'Excel cell text exceeds 32767 character limit','length':len(c.value)})
snap=values['情景快照'];count=0
for row in snap.iter_rows(min_row=7,values_only=True):
    if not row[0]:continue
    m=next(x for x in B['models'] if x['code']==row[0]);h=next(k for k,v in B['labels'].items() if v==row[2]);c=['悲观','基准','乐观'].index(row[3]);expected=m['scenarios'][h]['cases'][c]
    count+=1
    for i,k in [(5,'price'),(6,'dividend'),(7,'priceReturn'),(8,'totalReturn')]:
        ok=isinstance(row[i],(int,float)) and abs(row[i]-expected[k])<0.011 if i<=6 else isinstance(row[i],(int,float)) and abs(row[i]-expected[k])<1e-6
        if not ok:errors.append({'code':m['code'],'horizon':h,'case':c,'field':k,'actual':row[i],'expected':expected[k]})
    if row[19]!='已刷新':errors.append({'code':m['code'],'horizon':h,'status':row[19]})
year=values['年底情景'];matched=0
for row in year.iter_rows(min_row=7,max_row=109,values_only=True):
    m=next((x for x in B['models'] if x['code']==row[1]),None)
    if not m:continue
    matched+=1
    exp=m['scenarios']['yearEnd']['expectedPrice']
    if not isinstance(row[12],(int,float)) or abs(row[12]-exp)>.011:errors.append({'code':m['code'],'expectedPrice':exp,'actual':row[12]})
formulaCount=sum(1 for s in formulas for row in s for c in row if c.data_type=='f')
for sheet in formulas:
    for row in sheet:
        for cell in row:
            if cell.data_type=='f' and len(cell.value)>8192:errors.append({'sheet':sheet.title,'cell':cell.coordinate,'error':'Excel formula exceeds 8192 character limit','length':len(cell.value)})
for m in B.get('conditionedModels',[]):
    yr=next(row for row in year.iter_rows(min_row=7,max_row=109,values_only=True) if row[1]==m['code'])
    if any(v is not None for v in yr[6:17]):errors.append({'code':m['code'],'error':'Conditional reference entered formal year-end fields'})
    if m.get('basePriceValid') is False and yr[5] is not None:errors.append({'code':m['code'],'error':'Suspended historical price mislabelled valid base'})
    rows=[r for r in values['条件估值参考'].iter_rows(min_row=7,values_only=True) if r[2]==m['code']]
    if len(rows)!=len(m.get('conditionalScenarios',{})):errors.append({'code':m['code'],'error':'Conditional horizon coverage'})
    for row in rows:
        h=next(k for k,v in B['labels'].items() if v==row[8]);sc=m['conditionalScenarios'][h];prices=sc.get('priceHKD',sc.get('priceCNY',sc.get('priceTargets')))
        if any(not isinstance(row[10+c],(int,float)) or abs(row[10+c]-prices[c])>1e-8 for c in range(3)):errors.append({'code':m['code'],'horizon':h,'error':'Conditional source mismatch'})
report={'file':str(file),'sheets':values.sheetnames,'formulaCount':formulaCount,'scenarioRows':count,'endYearModels':matched,'errors':errors,'passed':not errors,'verification':'independent read-only openpyxl; no workbook authoring'}
(ROOT/'qa/workbook_qa.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='sheets'},ensure_ascii=False))
if errors:raise SystemExit(1)
