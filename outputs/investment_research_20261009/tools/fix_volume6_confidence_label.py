"""Translate one missing broker horizon label without changing research economics."""
import datetime, hashlib, json
from pathlib import Path
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
f=ROOT/'deliverables/06_公司深研_初版.docx'
d=Document(f)
old='中（公开转述，期限多数null）'
new='中（公开转述，目标期限多数未披露）'
cells=[c for t in d.tables for r in t.rows for c in r.cells if c.text==old]
assert len(cells)==1
before=hashlib.sha256(f.read_bytes()).hexdigest()
texts=lambda doc:[p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells]
original=texts(d)
count=0
for p in cells[0].paragraphs:
    for run in p.runs:
        if run.text=='中（公开转述，期限多数':
            run.text='中（公开转述，目标期限多数'
            count+=1
        elif run.text=='null':
            run.text='未披露'
            count+=1
assert count==2 and cells[0].text==new
d.save(f)
assert [t.replace(old,new) for t in original]==texts(Document(f))
q={'passed':True,'file':str(f),'beforeSHA256':before,'afterSHA256':hashlib.sha256(f.read_bytes()).hexdigest(),
   'createdAt':datetime.datetime.now(datetime.timezone.utc).isoformat(),'oldLabel':old,'newLabel':new,
   'change':'Presentation-only translation of literal null to missing target horizon; frozen inputs and economics unchanged.'}
(ROOT/'qa/volume6_confidence_label_fix.json').write_text(json.dumps(q,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(q,ensure_ascii=False))
