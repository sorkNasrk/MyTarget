"""Prevent the Kingdee bibliography link label splitting around a long URL."""
import datetime,hashlib,json
from pathlib import Path
from docx import Document
from docx.shared import Pt

ROOT=Path(__file__).resolve().parents[1]
f=ROOT/'deliverables/03_公司深研_初版.docx'
d=Document(f);before=hashlib.sha256(f.read_bytes()).hexdigest()
text=''.join(p.text for p in d.paragraphs)
p=[p for p in d.paragraphs if 'ew_INTERIMREPORT2026_1788424891861' in p.text]
assert len(p)==1
for run in p[0].runs: run.font.size=Pt(9.5)
d.save(f)
assert ''.join(p.text for p in Document(f).paragraphs)==text
q={'file':str(f),'beforeSHA256':before,'afterSHA256':hashlib.sha256(f.read_bytes()).hexdigest(),
   'createdAt':datetime.datetime.now(datetime.timezone.utc).isoformat(),'passed':True,
   'textUnchanged':True,'change':'Only Kingdee H1 interim source paragraph font9.5pt; full URL retained.'}
(ROOT/'qa/volume3_source_wrap_fix.json').write_text(json.dumps(q,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(q,ensure_ascii=False))
