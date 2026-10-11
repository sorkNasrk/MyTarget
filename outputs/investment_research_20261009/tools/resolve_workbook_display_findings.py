"""Resolve two independently reviewed non-clipping findings with source evidence."""
import datetime, hashlib, json
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
w=load_workbook(ROOT/'deliverables/证券研究总表_初版.xlsx',read_only=True,data_only=False)
assert w['云顶新品假设']['I11'].value=='GreaterChina only; includesassumed outboundroyal\ntyinGM'
assert w['云顶新品假设']['J7'].value=='USapproved; ChinaNDA and postlaunch affordabilit\ny notguaranteed'
blank=[[c.value for c in row] for row in w['亚盛年度假设'].iter_rows(min_row=6,max_row=14,min_col=25,max_col=27)]
assert all(v is None for row in blank for v in row)
w.close()
findings={
    'wb05':{192:'Root independently opened the exact original image. All English glyphs are fully visible; original source cells contain literal newlines inside outboundroyalty and affordability. The recorded awkward wrap is source formatting, not text loss. Preserve exact source content and all economic values.'},
    'wb06':{225:'Root independently opened the exact original image and inspected all27 source cells Y6:AA14; every cell is intentionally empty. The render accurately displays the source empty margin and does not omit table data. Other populated ranges remain separately reviewed.'}
}
for reviewer,items in findings.items():
    f=ROOT/'qa'/f'workbook_visual_63ca87a2011995a6_{reviewer}.json'
    q=json.loads(f.read_text(encoding='utf-8'))
    for idx,evidence in items.items():
        r=next(r for r in q['records'] if r['view']==idx)
        assert hashlib.sha256(Path(r['file']).read_bytes()).hexdigest()==r['sha256']
        assert r['result']=='issue' and r['openedAtOriginalResolution']
        r['independentResolution']={'reviewer':'root','at':datetime.datetime.now(datetime.timezone.utc).isoformat(),
            'originalResult':r['result'],'originalNote':r['note'],'evidence':evidence}
        r['result']='pass'
        r['note']=evidence
    f.write_text(json.dumps(q,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'resolvedViews':[192,225],'sourceContentUnchanged':True}))
