"""Compact the final Midea source paragraph; retain every character."""
import datetime, hashlib, json, re
from pathlib import Path
from docx import Document
from docx.shared import Pt

ROOT=Path(__file__).resolve().parents[1]
file=ROOT/'deliverables/02_公司深研_初版.docx'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
doc=Document(file)
before=sha(file)
text_before=''.join(p.text for p in doc.paragraphs)
matches=[p for p in doc.paragraphs if 'NOW.1527924' in p.text]
assert len(matches)==1
p=matches[0]
p.paragraph_format.line_spacing=1.1
p.paragraph_format.space_before=Pt(0)
p.paragraph_format.space_after=Pt(4)
for run in p.runs:
    run.font.size=Pt(9.5)
doc.save(file)
after=Document(file)
assert ''.join(p.text for p in after.paragraphs)==text_before
out={'createdAt':datetime.datetime.now(datetime.timezone.utc).isoformat(),'passed':True,
     'file':str(file),'beforeSHA256':before,'afterSHA256':sha(file),
     'paragraphTextUnchanged':True,'change':'Only final Midea bibliography paragraph: 9.5 pt, 1.1 line spacing, 4 pt after; all text and model values retained.'}
(ROOT/'qa/volume2_bibliography_fix.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(out,ensure_ascii=False))
