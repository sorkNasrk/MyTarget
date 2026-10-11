import argparse, json, math, re, os, sys, unicodedata
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent/'python_deps'))
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
import numpy as np

ROOT=Path(__file__).resolve().parents[1]
B=json.loads((ROOT/'data/research_bundle.json').read_text(encoding='utf-8'))
OUT=ROOT/'deliverables'
OUT.mkdir(exist_ok=True)
CHART=ROOT/'qa/charts'
CHART.mkdir(parents=True,exist_ok=True)
FONT=FontProperties(fname=r'C:\Windows\Fonts\msyh.ttc')
NAVY='24364B'
TEAL='31738A'
LABELS=B['labels']
ALL_MODELS=B['models']+B.get('conditionedModels',[])
HORIZONS=['halfYear','oneYear','threeYear','fiveYear']
DIMENSION_NAMES=['估值与预期回报','量价与相对强弱','经营增长与商业质量','行业周期与需求','期限催化剂','财务治理与下行风险']
DIMENSION_KEYS=['valuationExpectedReturn','technicalRelativeStrength','growthBusinessQuality','industryCycle','catalysts','financialGovernanceDownside']
DIMENSION_ALIASES=[['valuation','value','valuationReturn'],['technical','technicals','momentum'],['business','growth','quality','growthQuality'],['cycle','industry'],['catalyst'],['risk','financial','downside']]
VOLUME_SIZES=[8]*8+[7]*4
NETEASE_SCORE_TRANSLATIONS={
    'historical nonGAAP PE15.16x, cash strong but no double addition; Q2 investment and tax losses limit earnings certainty':'历史非公认会计准则市盈率约15.16倍；现金充足，估值不重复加回。第二季度投资与税项损失降低净利润的可预测性',
    'price above MA20/60/120 with 1.445 volume ratio, below MA250; weekly mixed; excess HSI20+8.66pp/60-2.19pp':'现价高于20、60和120日均线，量比1.445，仍低于250日均线；周线信号不一致。相对恒生指数，20日超额回报+8.66个百分点，60日为-2.19个百分点',
    'H1 game revenue+8.3%, game gross margin Q2 76.1%; mature franchises support operating income, content concentration matters':'上半年游戏收入增长8.3%，第二季度游戏毛利率76.1%；成熟游戏系列支撑经营利润，收入对重点内容的依赖仍需关注',
    'global content and AI efficiency opportunities, channel pricing and product supply competition constrain perpetuity':'全球内容需求和人工智能应用提供增长及效率机会；渠道议价和游戏供给竞争约束长期价值',
    'new and mature game update cycles and quarterly operating-to-net-profit bridge; unannounced release dates not invented':'新游戏推出、成熟游戏更新及季度经营利润到净利润的变化是观察催化；未公布的上线日期留空',
    'H1 OCF23.706bn, net cash167.5bn, light equipment spending; investments, taxes, SBC and cash deployment remain risks':'上半年经营现金流237.06亿元，净现金1,675亿元，设备支出较低；投资、税项、股份支付及现金配置仍有风险'
}

def display_confidence(value):
    """Translate evidence labels for presentation only; model data stays intact."""
    translations={'high':'高','medium':'中','low':'低'}
    if isinstance(value,str):
        return re.sub(r'\b(high|medium|low)\b',lambda match:translations[match.group(0).lower()],value,flags=re.I)
    if isinstance(value,dict):return {k:display_confidence(v) for k,v in value.items()}
    if isinstance(value,list):return [display_confidence(v) for v in value]
    return value

def display_evidence_field(value):
    """Translate evidence table keys without changing the underlying model."""
    return {'fundamentals':'基本面','primaryFundamentals':'基本面主源','technical':'量价','valuation':'估值','institutionTargets':'机构目标','brokerTargets':'机构目标'}.get(value,value)

def presentation_score_evidence(model,horizon):
    """Summarize existing issuer prose for display without altering model scores."""
    if model.get('issuerId')=='I006':
        h_share=model['code'].endswith('.HK')
        valuation={
            'halfYear':'按2025年每股盈利估算，A/H现价市盈率约56.3/44.6倍，已反映增长恢复；半年基准涨幅约6.22%/8.51%，追涨空间较有限',
            'oneYear':'一年基准每股盈利5.5元人民币，A/H合理倍数40/34倍；盈利改善需新兴业务、成熟项目和效率共同兑现',
            'threeYear':'三年基准每股盈利8.3元人民币，A/H倍数35/30倍；增长与新产能回报须兑现，不能把订单金额直接当利润',
            'fiveYear':'五年基准每股盈利11.5元人民币，A/H倍数32/28倍；成熟度提高后估值倍数下调，回报主要依靠经营扩张'
        }[horizon]+('。H股按研究汇率折算的现价较A股低，但价差不保证收敛' if h_share else '。A股现价估值高于H股，须有更强盈利兑现支撑')
        technical=('H股低于20日线而高于中长期线；基准日跌4.85%、量比1.58，放量下跌与周线升势并存' if h_share else 'A股略低20日线、高于中长期线；基准日量比约0.95，周线向上而短线动量回落') if horizon!='fiveYear' else '五年技术权重为零，量价不参与长期评分；两地短期波动和技术失效条件见第十二至十四节'
        return [valuation,technical,'上半年新兴业务增长72.49%，化学大分子与生物大分子扩张；集团归母利润下降15.71%，经营质量仍需交付与毛利验证','多肽、寡核苷酸及偶联外包有需求空间；客户融资、项目取消、价格和海外政策影响订单转化','后期项目交付、新产能利用率与后续商业化是主要催化；未披露客户药物及金额未另加价值','银行借款较少，但扩产支出超过经营现金流；库存、受限资金、激励及海外运营成本构成反证'], '凯莱英独立深研第六至十七节；以下为正文摘要，评分保持已核验模型原值'
    if model.get('issuerId')=='I082':
        valuation={
            'halfYear':'半年基准目标要求收入与正常利润率改善，现价尚未获得技术突破确认；回购、分红和股份变化分别计值',
            'oneYear':'一年基准收入5,300亿元、正常净利率9.9%、PE17倍；海外品牌与工业增量须抵消国内弱势',
            'threeYear':'三年基准收入6,200亿元、正常净利率10.2%、PE18倍；基准回报同时依赖经营、回购和合理倍数',
            'fiveYear':'五年基准收入7,250亿元、正常净利率10.5%、PE18倍；净回购与累计股息支持每股回报，未另加未经核验的投资资产价值'
        }[horizon]
        technical='五年技术权重为零，长期价值依赖经营与资本纪律，不能用短期技术反弹提高长期估值' if horizon=='fiveYear' else '现价低于20/60日线、高于120/250日线；日线上跌量比0.768，周线仍向上，短期调整未充分消化'
        return [valuation,technical,'规模、研发、供应链及全球品牌支撑经营；楼宇和机器人增长，工业零部件下滑，新增收入须有现金与利润','家庭耐用品更新、建筑节能和工业自动化提供持续需求；海外品牌与工业扩张受消费、贸易和资本开支周期约束','海外品牌利润、楼宇订单、机器人毛利和渠道改革效果需逐项验证；远期经营改善不视为已经兑现','经营现金较强但应付与预收不是永久净现金；跨国并购、激励、股数稀释及融资分红风险需持续核验'], '美的集团独立深研第五至十七节；以下为正文摘要，评分保持已核验模型原值'
    if model.get('issuerId')=='I066':
        valuation={
            'halfYear':'低预期修复提供基准上涨空间，但悲观损失深；估值需广告降幅收窄、云收入与回款改善支持',
            'oneYear':'正常利润依赖AI基础设施、应用与营销增量覆盖旧广告下降和共享费用；资本投入不能从利润估值中省略',
            'threeYear':'三年基准云基础设施规模扩大、应用收费提高并保留旧现金流，PE23倍要求合理贡献率和资本回报',
            'fiveYear':'五年基准云与应用规模有较高要求，PE23倍；现金收益已在EPS中，不再重复加净现金或子公司市值'
        }[horizon]
        technical='五年技术权重为零，资本回报与竞争决定长期价值；短期趋势不会直接提高远期目标' if horizon=='fiveYear' else '现价低于20/60/120/250日均线，斜率均负；20/60日相对恒生指数落后3.92/21.56个百分点，日周线尚弱'
        return [valuation,technical,'二季度AI基础设施增长而在线营销下降19%；五块收入来自管理账，独立利润率未披露，应用收费与设备回报待验证','企业AI、任务入口和自动驾驶有需求空间；云压价、搜索替代、芯片限制及设备迭代约束长期利润','云合同确认、应用续费、广告降幅收窄和资本节奏是验证项；双重主要上市已经生效，不再作为未来确定催化','二季度资本开支约为经营现金流3.3倍；现金回收、融资、可赎回少数权益与普通股稀释限制风险评分'], '百度独立深研第二至十三节；以下为正文摘要，评分保持已核验模型原值'
    if model.get('issuerId')=='I078':
        valuation={
            'halfYear':'半年基准依赖核心经营率恢复和市场提前承认正常利润，空间较大但图形未验证，悲观路径仍可大幅亏损',
            'oneYear':'一年基准核心收入与单位经济同时改善，新业务过渡投入必须受控；核心倍数是税后经营利润倍数',
            'threeYear':'三年基准核心收入4,500亿元、经营率17%，新业务质量与可分配资金改善；高基准回报需要竞争显著缓和',
            'fiveYear':'五年基准核心收入5,600亿元、经营率18%，核心倍数降至20倍；每股经营与现金增长须兑现，未另加战略投资市值'
        }[horizon]
        technical='五年技术权重为零，网络价值需跨周期现金与经营证明，日线反弹不单独提高长期价值' if horizon=='fiveYear' else '现价低于全部主要日均线，斜率均负；上跌量比0.74，20/60日相对恒生指数落后1.37/14.82个百分点'
        return [valuation,technical,'本地商户、履约和用户网络支持复购，2025年巨亏与2026半年现金减弱说明规模不保证利润恢复','本地服务与即时零售有空间；外卖补贴、到店竞争和海外本地化决定盈利周期，香港盈利不能外推全部海外','核心经营率穿越淡旺季、新业务亏损收敛及订单效率是催化；自动配送和AI必须有净成本节省','半年资本现金高于经营现金，股东现金代理为负；海外消耗、债务、骑手保障、信用及同股不同权限制下行保护'], '美团独立深研第一至五节及期限反证分析；以下为正文摘要，评分保持已核验模型原值'
    if model.get('issuerId')=='I070':
        valuation='正常盈利桥保留零售、物流、新业务亏损、常态费用与有限财资收益；'+{'halfYear':'半年回报依赖零售利润率与投入控制','oneYear':'一年需盈利改善与新业务亏损收窄','threeYear':'三年回报依赖经营率提升、物流利用率和每股增长','fiveYear':'五年仍保留新业务负贡献，回报不能只由收入规模推导'}[horizon]
        technical='五年技术权重为零，量价用于观察入场，不推导远期经营' if horizon=='fiveYear' else '现价略高20日线，仍低于60/120/250日线且斜率皆负；当日量比1.66，周线反转未确认'
        return [valuation,technical,'可信商品、供应链与交付网络支持业务；低毛利零售和新业务投入使利润对小幅经营率变化敏感','消费更新、品类扩展和企业服务有需求；家电高基数、外卖补贴与国际扩张使周期与资本投入承压','零售组合、物流外部利用率和新业务投入效率需后续业绩验证；消费活动未作未公布业绩预告','长期现金需覆盖物流、外卖和海外投入；激励抵消回购、子公司归属及持续新业务亏损可推翻每股增长假设'], '京东集团独立深研第一至五节；以下为正文摘要，评分保持已核验模型原值'
    if model.get('issuerId')=='I080':
        valuation='历史低PE不能替代目标时点普通股利润；'+{'halfYear':'半年依赖核心企稳与可灵收入的资金回报','oneYear':'一年需核心盈利和可灵投入共同改善','threeYear':'三年基准有较高增长与贡献率要求','fiveYear':'五年高空间需要核心现金、可灵持续商业化与资本纪律共同兑现'}[horizon]
        technical='五年技术权重为零，长期价值需真实客户现金与普通股归属支持' if horizon=='fiveYear' else '基准日涨5.47%、量比1.91，但价格仍低于四条下降日均线，周线偏弱；反弹尚未形成中期反转'
        return [valuation,technical,'成熟广告与社区支持现金基础，可灵二季度收入超过8.5亿元；集团利润与毛利下滑，收入增长尚未抵消AI投入','视频创作和商户工具有空间；注意力、广告、电商、模型价格与版权竞争限制长期贡献率','核心商家留存、可灵续费/API毛利与收入回款是验证项，模型更新不能替代盈利证据','半年资本及预付款高于经营现金，研究现金代理为负；子公司赎回权、激励、集团及可灵稀释需持续核验'], '快手独立深研第三至十四节；以下为正文摘要，评分保持已核验模型原值'
    if model.get('issuerId')=='I065':
        valuation='正常合并经营利润保留常态股份成本并使用经济摊薄股数；'+{'halfYear':'强季度已推动股价，半年仍需持续利润支持','oneYear':'一年需服务器高利润率持续和客户回款验证','threeYear':'三年基准依赖基础设施规模及每股正常盈利增长','fiveYear':'五年基准要求设备与服务资本回报穿越算力周期，未给现金重复加值'}[horizon]
        technical='五年技术权重为零，普通股价值需盈利、现金与资本结构支持' if horizon=='fiveYear' else '近期上升与经营改善提供动量，但价格已经大幅上行，量价及周线高波动要求后续回踩确认'
        return [valuation,technical,'PC供应链、服务器与服务能力有基础，首季服务器收入利润跃升；单季法定与调整利润差异大，不能把最佳季度永久化','AI基础设施和企业服务需求改善，PC仍有设备周期；OEM价格、芯片供应和客户资本开支限制利润持续性','后续服务器利润率、订单转收入、PC定价与服务现金是主要验证；pipeline不等于已签不可取消收入','营运资金快速占用，首季现金弱于调整利润；权证现金结算、可转债赎回、少数权益与经济股数稀释构成风险'], '联想集团独立深研第三至十七节；以下为正文摘要，评分保持已核验模型原值'
    return None,''

def coverage_state(bundle=B):
    summary=bundle['status']['summary']
    stocks=[s for s in bundle['universe']['securities'] if s['type']=='股票']
    issuer_count=len({s['issuerId'] for s in stocks})
    all_models=bundle['models']+bundle.get('conditionedModels',[])
    ready=summary.get('remainingDeepResearch',issuer_count)==0 and summary['deepResearchPassed']==issuer_count and len({m['issuerId'] for m in all_models})==issuer_count and len(all_models)==len(stocks)
    return {'complete':ready,'issuers':issuer_count,'securities':len(stocks),'summary':summary,'stocks':stocks}

def research_scope_text(bundle=B):
    state=coverage_state(bundle);ss=state['summary']
    if state['complete']:
        return f"研究覆盖{state['issuers']}家公司、{state['securities']}个股票证券，全部公司正文通过一万有效汉字门槛。{len(bundle['models'])}个可评估证券形成四期限排名，{len(bundle.get('conditionedModels',[]))}个条件研究证券保留独立分析并暂不排名。"
    return f"研究范围为{state['issuers']}家公司、{state['securities']}个股票证券。当前纳入{ss['issuerModels']}家公司模型，{ss['deepResearchPassed']}家正文通过一万有效汉字门槛，剩余{ss['remainingDeepResearch']}家公司深研待补齐。阶段排名仅纳入已完成且具备可评估条件的证券。"

def score_dimension_details(model,horizon):
    score=model.get('scores',{}).get(horizon,{})
    raw=score.get('explanations') or score.get('evidence') or score.get('reasons')
    if not raw:
        raw=model.get('scoreReasons',{})
        if isinstance(raw,dict) and horizon in raw:raw=raw[horizon]
    source_basis=''
    if not raw:raw,source_basis=presentation_score_evidence(model,horizon)
    rows=[]
    order=model.get('dimensionOrder') or DIMENSION_KEYS
    if not isinstance(order,list) or len(order)!=6:order=DIMENSION_KEYS
    for i,name in enumerate(DIMENSION_NAMES):
        reason=''
        item=raw[i] if isinstance(raw,list) and i<len(raw) else None
        if isinstance(raw,dict):item=next((raw[key] for key in [order[i],name]+DIMENSION_ALIASES[i] if key in raw),None)
        if isinstance(item,dict):reason=item.get('reason') or item.get('evidence') or item.get('explanation') or ''
        elif isinstance(item,str):reason=item
        if model.get('issuerId')=='I064':reason=NETEASE_SCORE_TRANSLATIONS.get(reason,reason)
        rows.append({'dimension':name,'score':score.get('dimensions',[None]*6)[i],'weight':B['weights'][horizon][i],'reason':reason,'sourceBasis':source_basis})
    return rows

def volume_groups(bundle=B):
    all_models=bundle['models']+bundle.get('conditionedModels',[])
    issuer_order=sorted({m['issuerId'] for m in all_models},key=lambda iid:(min((r['rank'] for rows in bundle['rankings'].values() for r in rows if r['issuerId']==iid),default=10000),-max((m.get('scores',{}).get('oneYear',{}).get('total',0) for m in all_models if m['issuerId']==iid),default=0),min(m['code'] for m in all_models if m['issuerId']==iid)))
    statuses={c['issuerId']:c for c in bundle['status']['companies']}
    complete=[iid for iid in issuer_order if statuses[iid]['effectiveHanCharacters']>=10000]
    groups=[];start=0
    for size in VOLUME_SIZES:
        group=complete[start:start+size];start+=size
        if group:groups.append(group)
    if coverage_state(bundle)['complete']:
        assert len(complete)==92 and [len(g) for g in groups]==VOLUME_SIZES,'完整公司卷须为8卷各8家公司及4卷各7家公司'
        assert len({iid for group in groups for iid in group})==92,'公司卷重复或遗漏'
    return groups

def stock_market_method():
    text=(ROOT/'research/market_method.md').read_text(encoding='utf-8')
    state=coverage_state();market={m['code']:m for m in B['market']['securities']}
    anchored=sum(market[s['code']].get('tradableOnAsOf',True) and market[s['code']].get('priceDate')==B['asOf'] for s in state['stocks'])
    old='103项均有真实可取得历史，其中100项有基准日价格锚，3项停牌仅保留最后交易参考价。'
    text=text.replace(old,f"{state['securities']}个股票证券均保留实际可取得历史，其中{anchored}个取得基准日有效价格锚，停牌证券仅保留最后交易参考价。")
    text=text.replace('A 股 ETF 日线可为手而 qt 仍为股；','')
    text=text.replace('部分基金 Yahoo 最新bar缺开盘价，追加腾讯真实完整bar，使用新浪股数核验最新成交量。','')
    text=text.replace('159502从2024-01-10上市首日，美的H股从2024-09-17公开上市日裁剪。07709各源开端不一致，在没有官方上市及初期数据核验前保留Yahoo实际可用225根，不补造早期日线。','美的H股从2024-09-17公开上市日裁剪。')
    # Index data is used only as a company relative-strength benchmark; no separate product research is delivered.
    text=re.sub(r'## 指数基准与适用范围\n.*?(?=## 目标交易日)', '## 相对强弱基准的适用范围\n\n恒生指数、沪深300及匹配的行业指数仅用于公司相对强弱比较。基准数据的来源、真实交易日期及量能限制保留在行情锁版；行业映射未核验时留空，不把粗粒度市场基准冒充公司的细分行业指数。\n\n',text,flags=re.S)
    return text

def pct(x): return '—' if x is None else f'{x:+.2%}'
def num(x): return '—' if x is None else f'{x:,.2f}'
def add_field(p,text):
    r=p.add_run(); f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),text);r._r.addnext(f)

def style_font(style,name='宋体',size=11):
    style.font.name=name;style.font.size=Pt(size)
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name)

def clean_headings(d):
    for paragraph in d.paragraphs:
        if paragraph.style.name=='Title' or paragraph.style.name.startswith('Heading'):
            text=''.join(' ' if unicodedata.category(c).startswith('P') else c for c in paragraph.text)
            paragraph.text=re.sub(r'\s+',' ',text).strip()
            for run in paragraph.runs:run.font.color.rgb=RGBColor(0,0,0)

def document(title,subtitle):
    d=Document();s=d.sections[0]
    s.page_width=Cm(21.59);s.page_height=Cm(27.94)
    s.top_margin=Cm(2);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2.1);s.right_margin=Cm(2.1)
    style_font(d.styles['Normal']);p=d.styles['Normal'].paragraph_format
    p.line_spacing=1.4;p.space_after=Pt(7);p.widow_control=True
    for name,size in [('Title',24),('Heading 1',17),('Heading 2',13),('Heading 3',11)]:
        style_font(d.styles[name],'微软雅黑',size);d.styles[name].font.color.rgb=RGBColor(0,0,0)
        pr=d.styles[name]._element.pPr
        if pr is not None:
            for border in list(pr.findall(qn('w:pBdr'))):pr.remove(border)
        d.styles[name].paragraph_format.keep_with_next=True
    style_font(d.styles['Caption'],'微软雅黑',9)
    s.header.paragraphs[0].text='医药与科技证券研究  |  价格基准 2026-10-09  |  初版'
    s.header.paragraphs[0].runs[0].font.size=Pt(8)
    foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run('研究假设可更新  ·  ');add_field(foot,'PAGE')
    d.core_properties.title=title;d.core_properties.subject='研究数据版本 '+B.get('fingerprint','未登记');d.core_properties.author='证券研究工作区'
    title=re.sub(r'[·—–:：，、（）()]',' ',title)
    title_p=d.add_paragraph(re.sub(r'\s+',' ',title).strip(),'Title');d.add_paragraph(re.sub(r'\s+',' ',''.join(' ' if unicodedata.category(c).startswith('P') else c for c in subtitle)).strip())
    for run in title_p.runs:run.font.color.rgb=RGBColor(0,0,0)
    status_text='全部公司深研和股票模型已覆盖，文档按同一数据版本生成。' if coverage_state()['complete'] else '公司深研卷收录已达到正文门槛的公司；剩余内容按研究状态继续补齐。'
    d.add_paragraph('资料截止：2026年10月9日。所有未来价格、盈利、估值、分红和汇率均为研究假设。'+research_scope_text()+status_text+'初版结论可随新增证据修订。数据版本 '+B.get('fingerprint','未登记')+'。')
    d.add_page_break()
    toc_title=d.add_paragraph('目录');toc_title.runs[0].bold=True;toc_title.runs[0].font.size=Pt(17)
    d.add_paragraph('[研究目录生成位置]')
    settings=d.settings.element
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    return d

def display_table_text(value):
    text=str(value)
    number=r'[+-]?\d[\d,.]*%?'
    if '.' in text and re.fullmatch(number+r'(?:/'+number+r'){1,}',text):
        return ' / '.join(text.split('/'))
    return text

def table(d,headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for i,h in enumerate(headers):
        t.rows[0].cells[i].text=str(h)
        tcpr=t.rows[0].cells[i]._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),NAVY);tcpr.append(shade)
        for r in t.rows[0].cells[i].paragraphs[0].runs:r.font.color.rgb=RGBColor(255,255,255);r.font.bold=True
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):cells[i].text=display_table_text(v)
    for ridx,row in enumerate(t.rows):
        keep=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(keep)
        for i,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx>0 and ridx%2==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F1F5F8');c._tc.get_or_add_tcPr().append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                v=OxmlElement('w:'+side);v.set(qn('w:w'),'90');v.set(qn('w:type'),'dxa');margins.append(v)
            c._tc.get_or_add_tcPr().append(margins)
            c.width=Cm(widths[i] if widths else 16.8/len(headers))
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.18
                word_wrap=OxmlElement('w:wordWrap');word_wrap.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(word_wrap)
                for r in p.runs:
                    r.font.name='微软雅黑';r.font.size=Pt(8.5);r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    d.add_paragraph()
    return t

def plot(m,weekly=False):
    hist=json.loads((ROOT/f"data/history/{m['securityId']}.json").read_text(encoding='utf-8'))
    fullbars=hist['weekly'] if weekly else hist['bars']
    bars=fullbars if weekly else fullbars[-180:]
    n=len(bars)
    fig,(ax,vol)=plt.subplots(2,1,figsize=(9,4.2),sharex=True,gridspec_kw={'height_ratios':[3,1]},layout='constrained')
    x=np.arange(n)
    c=np.array([b.get('adjusted',{}).get('close',b['close']) for b in bars])
    o=np.array([b.get('adjusted',{}).get('open',b['open']) for b in bars])
    h=np.array([b.get('adjusted',{}).get('high',b['high']) for b in bars])
    lo=np.array([b.get('adjusted',{}).get('low',b['low']) for b in bars])
    colors=np.where(c>=o,'#B85444','#3C8571')
    for j in range(n):
        ax.vlines(j,lo[j],h[j],color=colors[j],lw=.5)
        ax.bar(j,abs(c[j]-o[j]) or .001,bottom=min(c[j],o[j]),width=.65,color=colors[j],linewidth=0)
    fc=np.array([b.get('adjusted',{}).get('close',b['close']) for b in fullbars])
    for window,color in ([(10,'#8F6A20'),(30,'#406C91')] if weekly else [(20,'#8F6A20'),(60,'#406C91'),(120,'#875A85'),(250,'#486C56')]):
        ma=np.full(len(fullbars),np.nan)
        if len(fullbars)>=window:ma[window-1:]=np.convolve(fc,np.ones(window)/window,'valid')
        ma=ma[-n:]
        ax.plot(x,ma,color=color,lw=1,label=f'MA{window}')
    volume_ok=m.get('volumeChartsPermitted',True) and m.get('volumeMetricsUsable',True)
    if volume_ok:
        vv=np.array([b.get('volume',0) or 0 for b in bars])/1e6
        vol.bar(x,vv,color=colors,width=.7)
    else:
        vol.text(.5,.5,'成交量单位未完成交叉核验  暂不绘制量能',transform=vol.transAxes,ha='center',va='center',fontproperties=FONT,fontsize=9)
        vol.set_yticks([])
    dates=[b.get('date',b.get('weekEnd','')) for b in bars]
    ticks=np.linspace(0,n-1,min(6,n)).astype(int);vol.set_xticks(ticks,[dates[i][:10] for i in ticks],fontsize=8)
    ax.set_title(m['name']+(' — 周线' if weekly else ' — 最近180个交易日日线'),fontproperties=FONT,fontsize=12)
    ax.set_ylabel(m['currency'])
    if volume_ok:vol.set_ylabel('百万股',fontproperties=FONT,fontsize=8)
    ax.grid(axis='y',alpha=.15);vol.grid(axis='y',alpha=.15);ax.legend(loc='upper left',fontsize=8,ncol=3)
    for a in (ax,vol):a.spines[['top','right']].set_visible(False);a.tick_params(labelsize=8)
    dst=CHART/f"{m['securityId']}_{'weekly' if weekly else 'daily'}.png"
    fig.savefig(dst,dpi=160);plt.close(fig);return dst

def render_md(d,text):
    lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('# '):continue
        if line.startswith('|'):
            block=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i].strip());i+=1
            rows=[[c.strip().replace('**','') for c in r.strip('|').split('|')] for r in block if not re.fullmatch(r'[\s|:\-]+',r)]
            if rows:table(d,rows[0],rows[1:])
            continue
        if line.startswith('#'):
            d.add_paragraph(line.lstrip('# ').replace('**',''),'Heading 2');continue
        if re.match(r'^(价格基准|证券与价格基准|研究版本)',line):continue
        line=line.replace('**','').replace('`','')
        d.add_paragraph(line)

def conditional_presentation(m):
    if m.get('issuerId')=='I008':
        return (
            '以下是未融资完整计划的现金压力参考，展示现有资金与完整研发计划的资金缺口，不计算正式推荐分、涨跌幅、含息回报或判断加权期望价。远期资金缺口触发的0值不是近期期望价或清算价。',
            ['未融资计划期限','目标日','悲观现金压力参考','基准现金压力参考','乐观现金压力参考']
        )
    return (
        '以下仅是融资、存续或复牌条件成立后的估值参考，不计算正式推荐分、涨跌幅、含息回报或判断加权期望价。',
        ['条件参考期限','目标日','悲观参考价','基准参考价','乐观参考价']
    )

def model_summary(d,m):
    if m.get('rankEligible') is False:
        d.add_paragraph('暂不排名。'+m.get('rankingExclusionReason','核心数据或持续经营条件不足。'))
        conditional_note,conditional_headers=conditional_presentation(m)
        d.add_paragraph(conditional_note)
        table(d,conditional_headers,[[LABELS[h],s.get('target','待核')]+[num(v) for v in s.get('priceHKD',s.get('priceCNY',s.get('priceTargets',[None]*3)))] for h,s in m.get('conditionalScenarios',{}).items()])
        for h,s in m.get('conditionalScenarios',{}).items():
            d.add_paragraph(LABELS[h]+'成立条件：'+'；'.join(s.get('conditions',[])))
        return
    score=[]
    for k in ['halfYear','oneYear','threeYear','fiveYear']:
        r=next((r for r in B['rankings'][k] if r['code']==m['code']),None)
        score.append([LABELS[k],num(m['scores'][k]['total']),r['tier'] if r else '待全榜核验',num(m['scenarios'][k]['cases'][1]['price']),pct(m['scenarios'][k]['cases'][1]['priceReturn']),pct(m['scenarios'][k]['cases'][1]['totalReturn'])])
    table(d,['期限','得分/100','全榜梯队','基准目标价','股价涨幅','含分红回报'],score)
    comparison='梯队按全名单中具备合理评估条件的股票分别计算；条件研究证券明确暂不排名。' if coverage_state()['complete'] else '当前梯队按已完成且可评估证券计算；随着剩余公司模型接入，四榜统一刷新。'
    d.add_paragraph(comparison+'三年和五年年化以累计现金分红、不再投资计算。')
    s=m['scenarios']['yearEnd']
    table(d,['年底情景','目标价','涨跌幅','情景权重'],[[name,num(s['cases'][i]['price']),pct(s['cases'][i]['priceReturn']),w] for i,(name,w) in enumerate([('悲观','25%'),('基准','50%'),('乐观','25%')])])
    d.add_paragraph(f"年底判断加权期望价{num(s['expectedPrice'])}{m['currency']}，期望股价回报{pct(s['expectedPrice']/m['price']-1)}。权重是研究判断，并非统计概率；悲观价格亦不是最大损失边界。")
    table(d,['期限','悲观价格/涨幅','基准价格/涨幅','乐观价格/涨幅','基准年化总回报'],[[LABELS[h]]+[num(m['scenarios'][h]['cases'][c]['price'])+'\n'+pct(m['scenarios'][h]['cases'][c]['priceReturn']) for c in range(3)]+[pct(m['scenarios'][h]['cases'][1]['annualized'])] for h in ['halfYear','oneYear','threeYear','fiveYear']])
    table(d,['期限','名义目标日','模型采用日','日历核验状态'],[[LABELS[h],m['scenarios'][h].get('targetDateResolution',{}).get('nominalDate','待核'),m['scenarios'][h]['target'],m['scenarios'][h].get('targetDateResolution',{}).get('calendarStatus','待核')]for h in ['halfYear','oneYear','threeYear','fiveYear']],widths=[1.1,2.4,2.4,10.9])
    rebase=m['scenarios']['oneYear'].get('calendarRebase')
    if rebase:d.add_paragraph('一年日期修订：'+rebase['recalculation']+'。下文保留的额外压力测试按原假设日期运行，属于历史模型敏感性；核心三情景以上表及Excel当前快照为准。')
    confidence=display_confidence(m.get('confidence',{}))
    if confidence:table(d,['证据领域','可信度与限制'],[[display_evidence_field(k),v]for k,v in confidence.items()] if isinstance(confidence,dict) else [['总体',confidence]])
    if m['currency']=='HKD':
        rows=[]
        for h in ['halfYear','oneYear','threeYear','fiveYear']:
            total=m['scenarios'][h]['cases'][1]['totalReturn'];rows.append([LABELS[h],pct((1+total)*.95-1),pct(total),pct((1+total)*1.05-1)])
        table(d,['港币兑人民币变化','贬值5%总回报','不变总回报','升值5%总回报'],rows)
        d.add_paragraph('汇率敏感性固定本币目标价，仅转换人民币投资者回报；不会再把同一次汇率变动叠加到盈利换算假设。')

def overview():
    state=coverage_state()
    d=document('医药与科技股票研究 · 总报告','全名单结论、四期限排名、年底情景与研究方法' if state['complete'] else '股票名单总览、阶段四期限排名、年底情景与研究方法')
    d.add_paragraph('研究范围与初版口径','Heading 1').paragraph_format.page_break_before=True
    ss=B['status']['summary']
    d.add_paragraph(research_scope_text()+'研究交付包括股票公司深研和A/H证券差异；指数只作为相对强弱基准。')
    market_by_code={m['code']:m for m in B['market']['securities']}
    anchored=sum(market_by_code[s['code']].get('tradableOnAsOf',True) and market_by_code[s['code']].get('priceDate')==B['asOf'] for s in state['stocks'])
    history=sum((market_by_code[s['code']].get('historyCount') or 0)>0 for s in state['stocks'])
    d.add_paragraph(f"{anchored}个股票证券取得10月9日有效价格锚，{history}个保留实际可取得历史技术数据。微创脑科学、微创机器人和微创医疗自9月1日起停牌，使用8月31日最后交易参考价，暂不排名。其他条件研究证券的融资与存续限制单列，不以旧价计算正式期望回报。")
    d.add_paragraph('四期限推荐排序','Heading 1')
    for k in ['halfYear','oneYear','threeYear','fiveYear']:
        d.add_paragraph(LABELS[k]+(' · 全名单可评估股票' if state['complete'] else ' · 阶段可评估股票'),'Heading 2')
        leaders=[r for r in B['rankings'][k] if r['topTenDistinct']]
        table(d,['全榜名次','公司/优先代码','得分','基准目标','股价涨幅','悲观涨跌幅'],[[r['rank'],r['name']+'\n'+r['code'],num(r['score']),num(r['baseTarget']),pct(r['baseReturn']),pct(r['bearReturn'])] for r in leaders])
        d.add_paragraph('上表突出前十家不同公司，名次保留股票全榜位置；A/H其他代码与全部可评估证券见Excel。'+('全部公司研究已覆盖，暂不排名原因逐证券公开。' if state['complete'] else '阶段结果将在剩余公司接入后统一重排。'))
        for ranking in leaders:
            model=next(m for m in B['models'] if m['code']==ranking['code'])
            details=score_dimension_details(model,k)
            d.add_paragraph(model['name']+' '+model['code']+' 评分证据','Heading 3')
            d.add_paragraph('；'.join(item['dimension']+f" {item['score']}/5（权重{item['weight']}%）" for item in details)+'。')
            if details and details[0].get('sourceBasis'):d.add_paragraph(details[0]['sourceBasis']+'。')
            reasons=[item['dimension']+'：'+item['reason'] for item in details if item['reason']]
            if reasons:d.add_paragraph('；'.join(reasons)+'。')
            else:d.add_paragraph('结构化评分未单列简短证据说明；完整经营、估值、量价及反证分析见该公司独立深研，维度分值按已核验模型原值保留。')
    d.add_paragraph('进取型排序同时考虑上涨空间、经营兑现与期限催化；半年额外强调均线、量能和相对强弱，长期侧重商业质量和需求持续性。正常利润、授权及处置收益分开，创新药的适应症成功率、销售路径、现金投入及融资稀释逐公司建模。自有估值与机构目标存在分歧时，按假设及证据解释，不以机构目标直接替代年底情景。')
    d.add_paragraph('年底情景与期望值','Heading 1')
    table(d,['证券/币种','现价','悲观','基准','乐观','期望价','期望回报'],[[m['name']+'\n'+m['code']+' '+m['currency'],num(m['price']),num(m['scenarios']['yearEnd']['cases'][0]['price']),num(m['scenarios']['yearEnd']['cases'][1]['price']),num(m['scenarios']['yearEnd']['cases'][2]['price']),num(m['scenarios']['yearEnd']['expectedPrice']),pct(m['scenarios']['yearEnd']['expectedPrice']/m['price']-1)] for m in B['models']])
    d.add_paragraph('默认悲观/基准/乐观权重25%/50%/25%，另看40%/50%/10%与10%/50%/40%。12个月机构目标价只是估值参照，年底按截止信息集重新建模，没有按剩余月份线性缩放。')
    if B.get('conditionedModels'):
        d.add_paragraph('暂不排名公司的研究状态','Heading 2')
        table(d,['证券','参考报价/币种','参考日期','暂不排名原因'],[[m['name']+'\n'+m['code'],num(m.get('referencePrice',m['price']))+' '+m['currency'],m.get('referencePriceDate','待核'),m.get('rankingExclusionReason','待核')]for m in B['conditionedModels']])
        d.add_paragraph('上述公司有独立深研，条件估值详见深研和Excel专表；条件价格不与正常公司年底情景、正式回报或期望价混列。')
        if any(m.get('issuerId')=='I008' for m in B['conditionedModels']):
            d.add_paragraph('英矽智能暂不排名。其专表是未融资完整计划的现金压力参考；完整研发计划与现有资金不能在基准融资价格和项目退出假设下合理自洽。远期资金缺口触发的0值不是近期期望价、正式公允价或清算价，不纳入年底判断加权期望。完整独立研究保留经营、临床、技术和融资反证。')
    d.add_paragraph('评分与估值口径','Heading 1')
    d.add_paragraph('港股一年名义目标为2027年10月9日；10月8日为重阳节假期，且9日为周六，实际采用2027年10月7日。2026年12月31日为半日市，仍属有效交易日。香港政府2026及2027公布假期配合HKEX证券交易规则核验；2029和2031港股及内地远期日历未完成核验，保留名义日期而不填造实际日期。有限现金流模型已按已核验实际日重新折现，独立经营情景输入不按两日差异机械插值。')
    d.add_paragraph('日历主源：https://www.gov.hk/en/about/abouthk/holiday/2027.htm；https://www.hkex.com.hk/Services/Trading-hours-and-Severe-Weather-Arrangements/Trading-Hours/Securities-Market?sc_lang=en。公布月份分别为2026年5月与规则更新2017年9月16日。')
    table(d,['维度','半年','一年','三年','五年'],[[name]+[str(B['weights'][k][i])+'%' for k in ['halfYear','oneYear','threeYear','fiveYear']] for i,name in enumerate(['估值与预期收益','量价与相对强弱','经营与商业质量','周期与需求','催化剂','财务治理下行风险'])])
    d.add_paragraph('维度以0至5分按统一证据标尺判断：0为不可支持或极弱，1为明显不利，2为偏弱或关键证据不足，3为利弊较均衡，4为有多项可验证优势，5为优势突出且主要反证受控。加权得分不等于获利概率。每个期限分别评估，差距不足3分视为同梯队。')
    d.add_paragraph('成熟盈利公司用目标时点正常盈利及合理PE，现金流和资本配置交叉验证。亏损创新药区分商业化现金流、分适应症风险调整净现值和稀释，亏损科技保留收入毛利、完整投入和现金跑道；保险公司按承保、投资收益、净资产及PB交叉核验。A/H股共享发行人模型并独立计算证券估值和量价，无法合理评估时保留条件研究。')
    d.add_paragraph('行情与量价方法','Heading 1')
    render_md(d,stock_market_method())
    d.add_paragraph('经济、行业与收益条件','Heading 1')
    d.add_paragraph('创新药收益来自临床证据、商业化放量、授权和经营杠杆，资本市场融资改善需要经过研发投入、订单和审批才能进入现金流。授权总额含不确定里程碑，不能整体视为收入。医保可扩大可及人群，也会压低价格；海外审批、支付和专利竞争会使同一临床资产在不同地区的价值不同。CXO需求受全球药企预算和融资周期影响，订单总额、客户项目取消及海外产能资本开支决定现金回报。')
    d.add_paragraph('科技公司需要区别真实变现和主题预期。游戏与广告受消费和内容周期影响，云与AI受资本开支、推理价格和客户使用效益影响；制造与新能源汽车受价格竞争、产能利用率、海外关税和汇率影响。人民币与港币转换不仅改变人民币投资者回报，也影响以人民币报表盈利换算的港币目标价。因此本币结果与汇率敏感性分开展示，不假定汇率变化一定带来公司经营收益。')
    d.add_paragraph('全名单身份与状态','Heading 1')
    table(d,['证券标识','代码','名称','类别','模型/行情状态'],[[s['securityId'],s['code'],s['name'],s['type'],'情景研究完成' if any(m['code']==s['code'] for m in B['models']) else ('条件研究完成，暂不排名' if any(m['code']==s['code'] for m in B.get('conditionedModels',[])) else ('停牌，暂不排名' if not market_by_code[s['code']].get('tradableOnAsOf',True) else '公司研究待补齐'))] for s in state['stocks']],widths=[1.1,2.8,4.1,1.5,7.3])
    dest=OUT/'00_总报告_初版.docx';clean_headings(d);d.save(dest);return dest

def deep_volumes():
    outputs=[]
    for idx,group in enumerate(volume_groups(),1):
        subtitle=f"本卷{len(group)}家公司；按四榜最高名次、同名次一年得分及代码安排" if coverage_state()['complete'] else f"阶段收录{len(group)}家公司；剩余公司接入后按统一规则更新卷次"
        d=document(f'公司深研 · 第{idx:02d}卷',subtitle)
        for iid in group:
            listings=[m for m in ALL_MODELS if m['issuerId']==iid];m=listings[0]
            d.add_paragraph(m['name']+' · '+', '.join(x['code'] for x in listings),'Heading 1').paragraph_format.page_break_before=True
            for listing in listings:
                d.add_paragraph(listing['code']+(' 参考价 ' if listing.get('rankEligible') is False else ' 现价 ')+num(listing['price'])+' '+listing['currency'])
                model_summary(d,listing)
                for weekly in [False,True]:
                    d.add_picture(str(plot(listing,weekly)),width=Cm(16.6))
                    volume_note='成交量保留原股数。' if listing.get('volumeChartsPermitted',True) and listing.get('volumeMetricsUsable',True) else '成交量单位尚未完成交叉核验，量图和量能指标暂不展示。'
                    d.add_paragraph(('周线' if weekly else '日线')+'：真实OHLC按基准日尺度复权。'+volume_note+'来源及异常修复见行情锁版。'+('停牌公司仅绘最后真实交易日以前历史，不把停牌填充值作为交易。' if listing.get('basePriceValid') is False else ''), 'Caption')
            render_md(d,(ROOT/m['deepResearchFile']).read_text(encoding='utf-8'))
        dest=OUT/f'{idx:02d}_公司深研_初版.docx';clean_headings(d);d.save(dest);outputs.append(dest)
    return outputs

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--overview-only',action='store_true');args=parser.parse_args()
    files=[overview()]
    if not args.overview_only:files+=deep_volumes()
    (ROOT/'qa/document_manifest.json').write_text(json.dumps({'version':B['version'],'fingerprint':B.get('fingerprint'),'createdAt':B['createdAt'],'modelCount':len(B['models']),'completedIssuers':B['status']['summary']['deepResearchPassed'],'scope':'overview-only' if args.overview_only else 'all-current-volumes','status':'初版待逐页渲染验收','files':[str(p) for p in files]},ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'files':[str(p) for p in files]},ensure_ascii=False))
